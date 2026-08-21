"""
scl_agent.py
Watches a folder for SCL screenshots and appends each one to an Excel workbook.

A vision model reads the table instead of OCR. That is the whole reason this
file is short: local OCR needed grid detection, thresholding, per-cell
segmentation and page-mode tuning to read a table like this, and still made
mistakes. The model just reads it.

    pip install openpyxl
    set SCL_API_KEY=sk-or-...            (Windows;  export SCL_API_KEY=... elsewhere)
    python scl_agent.py

By default it watches the folder this file is sitting in, so dropping it into
your screenshots folder and running it is the whole setup. It remembers which
workbook you attached, in .scl_agent.json next to the images.

    python scl_agent.py --folder "C:\\shots" --excel "C:\\ref.xlsx"
    python scl_agent.py --once          process what is there, then exit
    python scl_agent.py --excel ...     change the workbook (remembered after)

What it does with each new image:
  - reads the 6 x 11 table (SX..SEQV for the six Membrane/Bending subtypes)
  - checks the numbers against identities that stress linearization guarantees
  - finds the last used row of the first worksheet, leaves one blank row, and
    writes the six rows beneath it, in columns C:M
  - logs the file to a Processing_History sheet, and skips anything already
    logged there, so re-running is safe
"""

import os
import io
import sys
import json
import math
import time
import base64
import hashlib
import argparse
import urllib.error
import urllib.request
from datetime import datetime

import openpyxl

# --- configuration ---------------------------------------------------------

API_URL = "https://openrouter.ai/api/v1/chat/completions"
MODEL = "anthropic/claude-opus-5"     # any vision model on OpenRouter works

STATE_FILE = ".scl_agent.json"        # lives in the watched folder
HISTORY_SHEET = "Processing_History"
POLL_SECONDS = 3

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}
MEDIA_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
               ".bmp": "image/bmp", ".tif": "image/tiff", ".tiff": "image/tiff"}

COLUMNS = ["SX", "SY", "SZ", "SXY", "SYZ", "SXZ", "S1", "S2", "S3", "SINT", "SEQV"]
SUBTYPES = ["Membrane", "Bending (Inside)", "Bending (Outside)",
            "Membrane+Bending (Inside)", "Membrane+Bending (Center)",
            "Membrane+Bending (Outside)"]

FIRST_COL = 3           # column C, unless the sheet's own headers say otherwise
HISTORY_HEADERS = ["Timestamp", "Image File", "Source Folder", "SCL Title",
                   "Rows Written", "Status", "Notes", "Fingerprint"]

PROMPT = f"""Read the stress-classification (SCL) results table in this screenshot.

It has exactly 6 data rows, in this order:
{chr(10).join(f"  {i+1}. {s}" for i, s in enumerate(SUBTYPES))}

and exactly 11 numeric columns, in this order:
  {", ".join(COLUMNS)}

Return ONLY a JSON object, no prose and no markdown fence:
{{"title": "<the SCL label, e.g. 'SCL- 10', or null>",
  "rows": [[11 numbers], [11 numbers], [11 numbers],
           [11 numbers], [11 numbers], [11 numbers]]}}

Rules:
- Copy each value EXACTLY as printed, including the minus sign and every
  decimal place. Do not round, reformat, or convert units.
- Ignore any chart, legend or axis labels below the table.
- If a cell is genuinely unreadable, use null for it. Never guess a digit.
"""


# --- reading the screenshot ------------------------------------------------

def api_key():
    for var in ("SCL_API_KEY", "OPENROUTER_API_KEY"):
        if os.environ.get(var):
            return os.environ[var]
    raise SystemExit(
        "No API key found. Get one at https://openrouter.ai/keys, then:\n"
        '  Windows:  set SCL_API_KEY=sk-or-...\n'
        '  Mac/Linux: export SCL_API_KEY=sk-or-...'
    )


def read_table(image_path, model=MODEL, timeout=120):
    """Ask the model for the table. Returns (title, rows) or raises RuntimeError."""
    with open(image_path, "rb") as f:
        raw = f.read()
    media = MEDIA_TYPES.get(os.path.splitext(image_path)[1].lower(), "image/png")
    data_url = f"data:{media};base64,{base64.b64encode(raw).decode()}"

    body = json.dumps({
        "model": model,
        "messages": [{"role": "user", "content": [
            {"type": "text", "text": PROMPT},
            {"type": "image_url", "image_url": {"url": data_url}},
        ]}],
        "response_format": {"type": "json_object"},
        "max_tokens": 2000,
    }).encode()

    req = urllib.request.Request(API_URL, data=body, method="POST", headers={
        "Authorization": f"Bearer {api_key()}",
        "Content-Type": "application/json",
        "X-Title": "SCL Extractor",
    })

    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        raise RuntimeError(f"API error {e.code}: {e.read().decode('utf-8', 'replace')[:300]}")
    except urllib.error.URLError as e:
        raise RuntimeError(f"Could not reach OpenRouter: {e.reason}")

    if "choices" not in payload:
        raise RuntimeError(f"Unexpected API response: {str(payload)[:300]}")

    text = payload["choices"][0]["message"]["content"].strip()
    if text.startswith("```"):        # some models fence their JSON anyway
        text = text.strip("`")
        text = text[text.find("{"):]
    try:
        parsed = json.loads(text)
    except ValueError:
        raise RuntimeError(f"Model did not return JSON: {text[:200]}")

    rows = parsed.get("rows")
    if not isinstance(rows, list) or len(rows) != 6:
        raise RuntimeError(f"Expected 6 rows, got {len(rows) if isinstance(rows, list) else type(rows).__name__}")
    for i, row in enumerate(rows):
        if not isinstance(row, list) or len(row) != 11:
            raise RuntimeError(f"Row {i+1} has {len(row) if isinstance(row, list) else '?'} values, expected 11")

    return parsed.get("title"), rows


# --- checking the numbers --------------------------------------------------

def validate(rows, rel_tol=1e-4):
    """
    Cross-check the block against identities that stress linearization
    guarantees, so a misread or invented digit is caught before it is written:

      1. Membrane == Membrane+Bending (Center)   -- bending is zero mid-surface
      2. Bending (Inside) == -Bending (Outside)  for SX..SXZ, being antisymmetric
      3. their S1/S2/S3 reverse and negate       -- negating a tensor does that
      4. their SINT and SEQV are identical       -- both invariant under negation

    Both reference samples satisfy all four exactly. rel_tol sits just above
    one unit in the last printed place (~3e-5), so genuine rounding passes and
    a wrong digit does not.
    """
    def close(a, b):
        if a is None or b is None:
            return False
        return math.isclose(a, b, rel_tol=rel_tol, abs_tol=1e-9)

    problems = []
    membrane, bend_in, bend_out, _mb_in, mb_centre, _mb_out = rows

    for j, name in enumerate(COLUMNS):
        if not close(membrane[j], mb_centre[j]):
            problems.append(f"{name}: Membrane ({membrane[j]}) != M+B Center ({mb_centre[j]})")

    for j in range(6):
        neg = -bend_out[j] if bend_out[j] is not None else None
        if not close(bend_in[j], neg):
            problems.append(f"{COLUMNS[j]}: Bending Inside ({bend_in[j]}) != -Outside ({bend_out[j]})")

    for j, k in ((6, 8), (7, 7), (8, 6)):   # S1<->-S3, S2<->-S2, S3<->-S1
        neg = -bend_out[k] if bend_out[k] is not None else None
        if not close(bend_in[j], neg):
            problems.append(f"{COLUMNS[j]}: Bending Inside ({bend_in[j]}) != -Outside {COLUMNS[k]} ({bend_out[k]})")

    for j in (9, 10):
        if not close(bend_in[j], bend_out[j]):
            problems.append(f"{COLUMNS[j]}: Bending Inside ({bend_in[j]}) != Outside ({bend_out[j]})")

    return problems


# --- the workbook ----------------------------------------------------------

def data_start_col(ws):
    """Where SX..SEQV actually begins in row 1; falls back to column C."""
    for start in range(1, 40):
        vals = [ws.cell(row=1, column=start + j).value for j in range(11)]
        if all(v is not None and str(v).strip().upper() == c for v, c in zip(vals, COLUMNS)):
            return start
    return FIRST_COL


def next_block_row(ws, col):
    """One blank row below the last row holding anything in the data columns."""
    for row in range(ws.max_row, 1, -1):
        if any(ws.cell(row=row, column=c).value not in (None, "") for c in range(col, col + 11)):
            return row + 2
    return 2        # empty sheet: start directly under the header


def log_history(wb, name, folder, title, rows_written, status, notes, fingerprint):
    if HISTORY_SHEET in wb.sheetnames:
        ws = wb[HISTORY_SHEET]
    else:
        ws = wb.create_sheet(HISTORY_SHEET)
        ws.append(HISTORY_HEADERS)
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)
    ws.append([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), name, folder,
               title or "", rows_written, status, notes, fingerprint])


def already_done(wb, fingerprint):
    if HISTORY_SHEET not in wb.sheetnames:
        return False
    for row in wb[HISTORY_SHEET].iter_rows(min_row=2, values_only=True):
        if row and len(row) >= 8 and row[7] == fingerprint and row[5] == "OK":
            return True
    return False


def fingerprint(path):
    with open(path, "rb") as f:
        data = f.read()
    return f"{os.path.basename(path)}:{len(data)}-{hashlib.md5(data).hexdigest()[:10]}"


def process(image_path, excel_path, model=MODEL):
    """Read one screenshot into the workbook. Returns a (status, message) pair."""
    name = os.path.basename(image_path)
    fp = fingerprint(image_path)

    try:
        wb = openpyxl.load_workbook(excel_path)
    except Exception as e:
        return "ERROR", f"cannot open workbook: {e}"

    if already_done(wb, fp):
        return "SKIPPED", "already in this workbook's history"

    try:
        title, rows = read_table(image_path, model)
    except RuntimeError as e:
        title, rows, problems = None, None, [str(e)]
    else:
        problems = validate(rows)

    if rows is None or problems:
        # Nothing is written when the numbers do not check out -- a flagged
        # image is cheap, a wrong stress value in the workbook is not.
        log_history(wb, name, os.path.dirname(image_path), title, "-",
                    "NEEDS REVIEW", "; ".join(problems)[:500], fp)
        _save(wb, excel_path)
        return "NEEDS REVIEW", "; ".join(problems)[:200]

    ws = wb[wb.sheetnames[0]]           # first worksheet only
    col = data_start_col(ws)
    start = next_block_row(ws, col)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            ws.cell(row=start + i, column=col + j, value=value)

    log_history(wb, name, os.path.dirname(image_path), title,
                f"{start}-{start + 5}", "OK", "Clean read.", fp)
    err = _save(wb, excel_path)
    if err:
        return "LOCKED", err
    return "OK", f"rows {start}-{start + 5}" + (f"  ({title})" if title else "")


def _save(wb, path):
    try:
        wb.save(path)
        return None
    except PermissionError:
        return f"{os.path.basename(path)} is open in Excel -- close it and it will retry"
    except Exception as e:
        return f"could not save: {e}"


# --- watching --------------------------------------------------------------

def load_state(folder):
    try:
        with open(os.path.join(folder, STATE_FILE), encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def save_state(folder, state):
    try:
        with open(os.path.join(folder, STATE_FILE), "w", encoding="utf-8") as f:
            json.dump(state, f, indent=2)
    except OSError:
        pass


def images_in(folder):
    try:
        return sorted(n for n in os.listdir(folder)
                      if os.path.splitext(n)[1].lower() in IMAGE_EXTS)
    except OSError:
        return []


def watch(folder, excel_path, model, once=False):
    seen = {}       # name -> (size, mtime) from the previous poll
    done = set()

    print(f"Watching : {folder}")
    print(f"Workbook : {excel_path}")
    print(f"Model    : {model}")
    print("Drop screenshots into the folder. Ctrl+C to stop.\n" if not once else "")

    while True:
        for name in images_in(folder):
            if name in done:
                continue
            path = os.path.join(folder, name)
            try:
                st = os.stat(path)
            except OSError:
                continue
            sig = (st.st_size, st.st_mtime)

            # A file still being written must not be read. Require its size and
            # timestamp to hold steady across two polls first.
            if seen.get(name) != sig:
                seen[name] = sig
                if not once:
                    continue
                time.sleep(1)

            print(f"  {name} ... ", end="", flush=True)
            status, message = process(path, excel_path, model)
            print(f"{status}  {message}")
            if status != "LOCKED":      # a locked workbook is retried next pass
                done.add(name)

        if once:
            return
        time.sleep(POLL_SECONDS)


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[2])
    ap.add_argument("--folder", default=here, help="folder to watch (default: this file's folder)")
    ap.add_argument("--excel", help="reference .xlsx (remembered after the first run)")
    ap.add_argument("--model", default=MODEL, help=f"OpenRouter model (default: {MODEL})")
    ap.add_argument("--once", action="store_true", help="process what is there, then exit")
    args = ap.parse_args()

    folder = os.path.abspath(args.folder)
    if not os.path.isdir(folder):
        raise SystemExit(f"Not a folder: {folder}")

    state = load_state(folder)
    excel = args.excel or state.get("excel_path", "")

    # Nothing attached and exactly one workbook sitting here: use it.
    if not excel:
        candidates = [n for n in os.listdir(folder)
                      if n.lower().endswith(".xlsx") and not n.startswith("~$")]
        if len(candidates) == 1:
            excel = os.path.join(folder, candidates[0])
            print(f"Using the only workbook in this folder: {candidates[0]}\n")

    if not excel:
        raise SystemExit(
            "No reference workbook. Pass one once and it will be remembered:\n"
            '  python scl_agent.py --excel "C:\\path\\to\\reference.xlsx"'
        )

    excel = os.path.abspath(excel)
    if not os.path.isfile(excel):
        raise SystemExit(f"No workbook at {excel}")

    if state.get("excel_path") != excel:
        save_state(folder, {"excel_path": excel})

    api_key()   # fail now with a clear message, not on the first screenshot

    try:
        watch(folder, excel, args.model, once=args.once)
    except KeyboardInterrupt:
        print("\nStopped.")


if __name__ == "__main__":
    main()
