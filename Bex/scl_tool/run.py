"""
run.py
Interactive entry point for the SCL Excel Extractor.

Usage:
    python run.py                  -> opens the interactive menu
    python run.py process          -> processes new images right away
    python run.py set-folder PATH  -> sets the reference image folder
    python run.py set-excel PATH   -> sets the reference Excel file
    python run.py history          -> prints a summary of the history log
"""

import os
import sys

import scl_core as core

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_PATH = os.path.join(SCRIPT_DIR, "config.json")


def cmd_process():
    cfg = core.load_config(CONFIG_PATH)
    if not cfg["image_folder"] or not cfg["excel_path"]:
        print("Reference folder and/or Excel file are not set yet.")
        print("Use option 2/3 in the menu, or run:")
        print("  python run.py set-folder \"C:\\path\\to\\images\"")
        print("  python run.py set-excel \"C:\\path\\to\\reference.xlsx\"")
        return

    print(f"Image folder : {cfg['image_folder']}")
    print(f"Excel file   : {cfg['excel_path']}")
    print("Scanning for new images...")
    try:
        summary = core.process_folder(cfg)
    except Exception as e:
        print(f"ERROR: {e}")
        return

    if summary["header_warnings"]:
        print("\nHeader check warnings (columns C1:M1 didn't fully match the expected layout):")
        for w in summary["header_warnings"]:
            print(f"  - {w}")

    print(f"\nProcessed : {len(summary['processed'])}")
    for fname, start, end in summary["processed"]:
        print(f"  - {fname}  ->  rows {start}-{end}")

    print(f"Skipped (already done): {len(summary['skipped'])}")
    for fname in summary["skipped"]:
        print(f"  - {fname}")

    print(f"Flagged for review (couldn't read reliably): {len(summary['flagged'])}")
    for fname, warnings in summary["flagged"]:
        print(f"  - {fname}")
        for w in warnings:
            print(f"      {w}")

    print("\nAll results (including flagged files) are logged in the "
          "'Processing_History' tab of the Excel file.")


def cmd_set_folder(path=None):
    cfg = core.load_config(CONFIG_PATH)
    if path is None:
        path = input("New image folder path: ").strip().strip('"')
    if not os.path.isdir(path):
        print(f"WARNING: '{path}' doesn't exist (saving anyway).")
    cfg["image_folder"] = path
    core.save_config(CONFIG_PATH, cfg)
    print(f"Image folder set to: {path}")


def cmd_set_excel(path=None):
    cfg = core.load_config(CONFIG_PATH)
    if path is None:
        path = input("New reference Excel file path: ").strip().strip('"')
    if not os.path.isfile(path):
        print(f"WARNING: '{path}' doesn't exist (saving anyway).")
    cfg["excel_path"] = path
    core.save_config(CONFIG_PATH, cfg)
    print(f"Reference Excel set to: {path}")


def cmd_history():
    cfg = core.load_config(CONFIG_PATH)
    if not cfg["excel_path"] or not os.path.isfile(cfg["excel_path"]):
        print("No reference Excel file is set yet.")
        return
    import openpyxl
    wb = openpyxl.load_workbook(cfg["excel_path"], data_only=True)
    if core.HISTORY_SHEET_NAME not in wb.sheetnames:
        print("No history recorded yet.")
        return
    ws = wb[core.HISTORY_SHEET_NAME]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) <= 1:
        print("No history recorded yet.")
        return
    print(f"{'Timestamp':<20} {'File':<25} {'Status':<12} {'Rows':<10} Notes")
    print("-" * 100)
    for row in rows[1:]:
        ts, fname, folder, title, sheet, rng, status, notes = (list(row) + [""] * 8)[:8]
        print(f"{str(ts):<20} {str(fname):<25} {str(status):<12} {str(rng):<10} {notes}")


def cmd_set_tesseract(path=None):
    cfg = core.load_config(CONFIG_PATH)
    if path is None:
        path = input("Path to tesseract.exe (leave blank to clear override): ").strip().strip('"')
    cfg["tesseract_cmd"] = path
    core.save_config(CONFIG_PATH, cfg)
    print("Tesseract path override updated.")


def show_status():
    cfg = core.load_config(CONFIG_PATH)
    print("Current settings:")
    print(f"  Image folder : {cfg['image_folder'] or '(not set)'}")
    print(f"  Excel file   : {cfg['excel_path'] or '(not set)'}")
    if cfg.get("tesseract_cmd"):
        print(f"  Tesseract    : {cfg['tesseract_cmd']}")


def menu():
    while True:
        print("\n=== SCL Excel Extractor ===")
        show_status()
        print("""
1) Process new images now
2) Change reference image folder
3) Change reference Excel file
4) View processing history
5) Set custom Tesseract-OCR path (only needed if not on your system PATH)
6) Exit
""")
        choice = input("Choose an option (1-6): ").strip()
        if choice == "1":
            cmd_process()
        elif choice == "2":
            cmd_set_folder()
        elif choice == "3":
            cmd_set_excel()
        elif choice == "4":
            cmd_history()
        elif choice == "5":
            cmd_set_tesseract()
        elif choice == "6":
            break
        else:
            print("Please enter a number from 1 to 6.")


if __name__ == "__main__":
    args = sys.argv[1:]
    if not args:
        menu()
    elif args[0] == "process":
        cmd_process()
    elif args[0] == "set-folder":
        cmd_set_folder(args[1] if len(args) > 1 else None)
    elif args[0] == "set-excel":
        cmd_set_excel(args[1] if len(args) > 1 else None)
    elif args[0] == "history":
        cmd_history()
    elif args[0] == "set-tesseract":
        cmd_set_tesseract(args[1] if len(args) > 1 else None)
    else:
        print(__doc__)
