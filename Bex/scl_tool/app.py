"""
app.py
Streamlit UI for the SCL Excel Extractor.

Run locally with:
    streamlit run app.py

Two ways to use it:
  - Upload mode (works anywhere, including when deployed to Streamlit
    Community Cloud): upload the reference Excel + one or more screenshots,
    process, then download the updated workbook.
  - Local folder mode (only works when running on your own machine, since
    it reads/writes files directly on disk): point at a folder and an
    Excel file by path, same as the command-line tool.
"""

import io
import os
import json

import openpyxl
import pandas as pd
import streamlit as st

import scl_core as core

st.set_page_config(page_title="SCL Excel Extractor", page_icon="\U0001F4CA", layout="wide")

CONFIG_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config.json")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def get_history_df(wb):
    if core.HISTORY_SHEET_NAME not in wb.sheetnames:
        return pd.DataFrame(columns=core.HISTORY_HEADERS)
    ws = wb[core.HISTORY_SHEET_NAME]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    return pd.DataFrame(rows, columns=core.HISTORY_HEADERS)


def show_summary(summary):
    c1, c2, c3 = st.columns(3)
    c1.metric("Processed", len(summary["processed"]))
    c2.metric("Skipped (already done)", len(summary["skipped"]))
    c3.metric("Flagged for review", len(summary["flagged"]))

    if summary["header_warnings"]:
        with st.expander("\u26A0\uFE0F Header check warnings", expanded=False):
            for w in summary["header_warnings"]:
                st.write(f"- {w}")

    if summary["processed"]:
        st.success(
            "Written to the workbook:\n\n"
            + "\n".join(f"- **{f}** \u2192 rows {s}-{e}" for f, s, e in summary["processed"])
        )
    if summary["skipped"]:
        st.info("Already processed previously, skipped:\n\n" + "\n".join(f"- {f}" for f in summary["skipped"]))
    if summary["flagged"]:
        st.warning("Could not confidently read these — logged as NEEDS REVIEW, nothing written:")
        for fname, warnings in summary["flagged"]:
            st.write(f"**{fname}**")
            for w in warnings:
                st.caption(f"  {w}")


# ---------------------------------------------------------------------------
# UI
# ---------------------------------------------------------------------------

st.title("SCL Excel Extractor")
st.caption(
    "Reads SCL stress-classification screenshots and fills columns C:M of the "
    "first worksheet in your reference workbook, keeping a history of every file processed."
)

mode = st.radio(
    "Mode",
    ["Upload files", "Local folder (this machine only)"],
    horizontal=True,
    help="'Upload files' works anywhere, including a hosted deployment. "
         "'Local folder' reads/writes directly on disk and only works when "
         "you're running this app on your own computer.",
)

st.divider()

# ---------------------------------------------------------------------------
# Upload mode
# ---------------------------------------------------------------------------
if mode == "Upload files":
    col1, col2 = st.columns(2)
    with col1:
        excel_file = st.file_uploader("Reference Excel workbook", type=["xlsx"])
    with col2:
        image_files = st.file_uploader(
            "Screenshot image(s)",
            type=["png", "jpg", "jpeg", "bmp", "tif", "tiff"],
            accept_multiple_files=True,
        )

    if excel_file is not None:
        key = f"{excel_file.name}-{excel_file.size}"
        if st.session_state.get("wb_key") != key:
            st.session_state["wb_key"] = key
            st.session_state["wb_bytes"] = excel_file.getvalue()
            st.session_state["wb_name"] = excel_file.name

    has_wb = "wb_bytes" in st.session_state

    process_clicked = st.button(
        "Process images",
        type="primary",
        disabled=not (has_wb and image_files),
    )

    if process_clicked:
        wb = openpyxl.load_workbook(io.BytesIO(st.session_state["wb_bytes"]))
        image_items = [{"name": f.name, "data": f.getvalue()} for f in image_files]
        with st.spinner("Reading images and updating the workbook..."):
            summary = core.process_workbook(wb, image_items, source_label="Uploaded via Streamlit")
        buf = io.BytesIO()
        wb.save(buf)
        st.session_state["wb_bytes"] = buf.getvalue()
        st.session_state["last_summary"] = summary

    if "last_summary" in st.session_state:
        show_summary(st.session_state["last_summary"])

    if has_wb:
        st.download_button(
            "Download updated workbook",
            data=st.session_state["wb_bytes"],
            file_name=st.session_state.get("wb_name", "updated.xlsx"),
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        st.caption(
            "Any formulas elsewhere in the workbook that depend on this data "
            "(other tabs, or other columns on this sheet) will recalculate "
            "automatically the next time you open the file in Excel."
        )

        with st.expander("Processing history", expanded=False):
            wb_preview = openpyxl.load_workbook(io.BytesIO(st.session_state["wb_bytes"]), data_only=True)
            df = get_history_df(wb_preview)
            if df.empty:
                st.write("No history yet.")
            else:
                st.dataframe(df, use_container_width=True)
    else:
        st.info("Upload a reference Excel workbook to get started.")

# ---------------------------------------------------------------------------
# Local folder mode
# ---------------------------------------------------------------------------
else:
    cfg = core.load_config(CONFIG_PATH)

    col1, col2 = st.columns(2)
    with col1:
        folder = st.text_input("Image folder path", value=cfg.get("image_folder", ""))
    with col2:
        excel_path = st.text_input("Reference Excel file path", value=cfg.get("excel_path", ""))

    tess_path = st.text_input(
        "Tesseract-OCR path override (optional, only if not on your system PATH)",
        value=cfg.get("tesseract_cmd", ""),
    )

    if st.button("Save settings"):
        cfg["image_folder"] = folder
        cfg["excel_path"] = excel_path
        cfg["tesseract_cmd"] = tess_path
        core.save_config(CONFIG_PATH, cfg)
        st.success("Settings saved.")

    st.divider()

    process_clicked = st.button("Process new images now", type="primary")

    if process_clicked:
        cfg["image_folder"] = folder
        cfg["excel_path"] = excel_path
        cfg["tesseract_cmd"] = tess_path
        try:
            with st.spinner("Scanning folder and updating the workbook..."):
                summary = core.process_folder(cfg)
            st.session_state["last_summary_local"] = summary
        except Exception as e:
            st.error(str(e))

    if "last_summary_local" in st.session_state:
        show_summary(st.session_state["last_summary_local"])

    if excel_path and os.path.isfile(excel_path):
        with st.expander("Processing history", expanded=False):
            wb_preview = openpyxl.load_workbook(excel_path, data_only=True)
            df = get_history_df(wb_preview)
            if df.empty:
                st.write("No history yet.")
            else:
                st.dataframe(df, use_container_width=True)
