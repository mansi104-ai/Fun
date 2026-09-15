"""
Writes extracted blocks into columns C:M of a workbook's first worksheet, in
place, at the path the workbook already lives at.

The layout is the one the reference workbook already uses: row 1 is the header
(C1:M1 = SX ... SEQV), then blocks of six rows with a single blank row between
them. Column A carries the SCL number and column B the row labels; neither is
touched, because the numbers are all that is being added.

"After the last occupied row" means the last row holding anything in C:M --
not the last row on the sheet. That distinction matters for the real workbook:
its column A is pre-numbered with empty blocks far below where the data ends
(numbers at rows 37, 44, 51, ... while C:M stops at row 35). Measuring the
whole sheet would push every new block hundreds of rows down; measuring C:M
puts it at row 37, landing exactly on the next pre-numbered block.
"""

import math
import shutil
from datetime import datetime

import openpyxl

FIRST_COL = 3          # column C
COL_COUNT = 11         # C..M inclusive
LAST_COL = FIRST_COL + COL_COUNT - 1
ROWS_PER_BLOCK = 6


def as_number(value):
    """
    Parse one reviewed cell into a float, or None if it isn't a number.

    An emptied cell comes back from the review table as NaN (or the text
    "nan"), not None. NaN is a float, so without the explicit check it would
    count as a value and be written into the sheet as a number that isn't one.
    """
    if value is None:
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value).strip().replace(",", "")
        # A minus sign pasted from elsewhere is often an en-dash; keep the
        # sign rather than losing the value to it.
        text = text.replace("–", "-").replace("—", "-").replace("−", "-")
        try:
            number = float(text)
        except ValueError:
            return None
    return number if math.isfinite(number) else None


def next_block_row(worksheet):
    """One blank row below the last row with anything in C:M; row 2 if empty."""
    for row in range(worksheet.max_row, 1, -1):
        if any(worksheet.cell(row=row, column=c).value not in (None, "")
               for c in range(FIRST_COL, LAST_COL + 1)):
            return row + 2
    return 2


def check_block(rows):
    """Return a reason the block can't be written, or None if it's good."""
    if len(rows) != ROWS_PER_BLOCK:
        return f"{len(rows)} rows, expected {ROWS_PER_BLOCK}"
    for i, row in enumerate(rows, start=1):
        values = [as_number(v) for v in row]
        if len(values) != COL_COUNT or any(v is None for v in values):
            missing = sum(1 for v in values if v is None)
            return f"row {i} has {COL_COUNT - missing} of {COL_COUNT} numbers"
    return None


def back_up(excel_path):
    """
    Copy the workbook beside itself before the first write of a save.

    openpyxl rewrites the whole file from its parsed model rather than patching
    the bytes, so anything it does not model is not carried across. This
    workbook holds only sheets, values and merged cells, all of which survive
    -- but a backup costs a moment and makes a bad save undoable.
    """
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = excel_path.with_name(f"{excel_path.stem}.backup-{stamp}{excel_path.suffix}")
    shutil.copy2(excel_path, backup)
    return backup


def append_blocks(excel_path, blocks, make_backup=True):
    """
    Append each block into C:M of the first worksheet and save to the same path.

    `blocks` is a sequence of (name, rows). A block that isn't exactly 6 x 11
    numbers is skipped whole and reported -- a half-written block on the sheet
    would be worse than none. Returns (written, skipped, backup) where written
    is [(name, first_row, last_row)] and skipped is [(name, reason)].
    """
    good, skipped = [], []
    for name, rows in blocks:
        reason = check_block(rows)
        if reason:
            skipped.append((name, reason))
        else:
            good.append((name, [[as_number(v) for v in row] for row in rows]))

    if not good:
        return [], skipped, None

    backup = back_up(excel_path) if make_backup else None

    workbook = openpyxl.load_workbook(excel_path)
    worksheet = workbook[workbook.sheetnames[0]]

    written = []
    for name, rows in good:
        start = next_block_row(worksheet)
        for i, values in enumerate(rows):
            for j, value in enumerate(values):
                cell = worksheet.cell(row=start + i, column=FIRST_COL + j, value=value)
                # Some empty cells in the reference workbook still carry a
                # 0.00E+00 format, left behind by earlier pasting, which
                # would show -0.040742 as -4.07E-02. General shows the whole
                # value as read.
                cell.number_format = "General"
        written.append((name, start, start + ROWS_PER_BLOCK - 1))

    workbook.save(excel_path)
    return written, skipped, backup


def describe_target(excel_path):
    """(sheet name, next row a block would land on) for the first worksheet."""
    workbook = openpyxl.load_workbook(excel_path, read_only=False)
    worksheet = workbook[workbook.sheetnames[0]]
    return worksheet.title, next_block_row(worksheet)
